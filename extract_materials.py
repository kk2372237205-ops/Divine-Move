from __future__ import annotations

import json
from pathlib import Path

from docx import Document


ROOT = Path(r"C:\Users\23722\Desktop\五子棋机器人")
SOURCE = ROOT / "_materials_readonly" / "下棋机械臂资源-0701"
OUT = ROOT / "_analysis"


def extract_docx(path: Path) -> dict:
    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)
    image_parts = [part for part in document.part.related_parts.values() if part.content_type.startswith("image/")]
    return {
        "file": str(path.relative_to(SOURCE)),
        "paragraphs": paragraphs,
        "tables": tables,
        "image_count": len(image_parts),
        "sections": len(document.sections),
    }


def main() -> None:
    OUT.mkdir(exist_ok=True)
    records = [extract_docx(path) for path in sorted(SOURCE.rglob("*.docx"))]
    (OUT / "tutorials_extracted.json").write_text(
        json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    chunks = ["# 下棋机械臂资源教程全文提取", ""]
    for record in records:
        chunks.extend(
            [
                f"## {record['file']}",
                f"内嵌图片：{record['image_count']}，文档分节：{record['sections']}",
                "",
                "### 正文",
            ]
        )
        chunks.extend(record["paragraphs"])
        if record["tables"]:
            chunks.extend(["", "### 表格"])
            for table_index, table in enumerate(record["tables"], start=1):
                chunks.append(f"表 {table_index}")
                for row in table:
                    chunks.append(" | ".join(row))
        chunks.append("")
    (OUT / "tutorials_extracted.md").write_text("\n".join(chunks), encoding="utf-8")

    for source in sorted(SOURCE.rglob("*.py")):
        destination = OUT / f"{source.parent.name}_{source.name}.txt"
        destination.write_text(source.read_text(encoding="utf-8"), encoding="utf-8")


if __name__ == "__main__":
    main()
